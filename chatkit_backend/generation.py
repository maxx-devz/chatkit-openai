"""Bounded document/image tools. Scope and permissions never come from model arguments."""
import base64
import io
import os
import re
from uuid import uuid4
from pathlib import Path

from agents import function_tool
from docx import Document
from docx.shared import Pt
from pydantic import BaseModel, Field
from chatkit.widgets import WidgetTemplate
from ai_provider import has_openai_hosted_key, openai_hosted_client
from provider_errors import ImageGenerationError, image_generation_error, image_provider_error

FILE_CARD = WidgetTemplate.from_file(str(Path(__file__).with_name("generated-file.widget")))
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def tool_failure(_context, error):
    if isinstance(error, ValueError):
        return str(error)
    return "The file could not be generated. Please retry or contact staff to check the generation service."


class DocumentSection(BaseModel):
    heading: str = Field(max_length=200)
    paragraphs: list[str] = Field(max_length=30)
    bullets: list[str] = Field(max_length=30)


def render_document(title, sections, company):
    if not title.strip() or len(title) > 200 or not 1 <= len(sections) <= 20:
        raise ValueError("Use a title and 1-20 document sections")
    if sum(len(part) for section in sections for part in [section.heading, *section.paragraphs, *section.bullets]) > 30000:
        raise ValueError("Document content exceeds 30000 characters")
    document = Document()
    document.core_properties.title = title
    document.core_properties.author = company
    document.styles["Normal"].font.name = "Calibri"
    document.styles["Normal"].font.size = Pt(11)
    document.add_heading(title, 0)
    for section in sections:
        if section.heading:
            document.add_heading(section.heading, 1)
        for paragraph in section.paragraphs:
            document.add_paragraph(paragraph)
        for bullet in section.bullets:
            document.add_paragraph(bullet, style="List Bullet")
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


async def reserve_tool(context, kind):
    if not context.config.get(kind):
        raise ValueError("This tool is not enabled for this client")
    context.tool_calls += 1
    if context.tool_calls > 2:
        raise ValueError("At most two files can be generated in one reply")
    if context.preview:
        return  # The preview endpoint enforces a separate staff hourly allowance.
    limit = context.config["documentLimit" if kind == "documents" else "imageLimit"]
    cursor = await context.db.execute("""
        INSERT INTO portal_assistant_tool_usage(client_id,kind,period_start,attempts)
        VALUES(%s,%s,DATE_TRUNC('month',NOW() AT TIME ZONE 'UTC')::date,1)
        ON CONFLICT(client_id,kind,period_start) DO UPDATE
        SET attempts=portal_assistant_tool_usage.attempts+1
        WHERE portal_assistant_tool_usage.attempts < %s RETURNING attempts
    """, (context.client_id, kind, limit))
    if not await cursor.fetchone():
        raise ValueError("The client's monthly generation allowance has been reached")


async def save_artifact(context, title, extension, mime, content, emit=None):
    if not 0 < len(content) <= 3000000:
        raise ValueError("Generated file exceeds the 3 MB download limit; request a smaller output")
    artifact_id = str(uuid4())
    stem = re.sub(r"[^a-zA-Z0-9 -]", "", title).strip()[:70] or "generated-file"
    filename = stem + extension
    await context.db.execute("""
        INSERT INTO portal_assistant_files(id,client_id,user_id,thread_id,preview,filename,mime_type,content)
        VALUES(%s,%s,%s,%s,%s,%s,%s,%s)
    """, (artifact_id, context.client_id, context.user_id, context.thread_id, context.preview, filename, mime, content))
    url = context.portal_origin + "/api/assistant-files/" + artifact_id + ("?preview=1" if context.preview else "")
    artifact = {"id": artifact_id, "filename": filename, "url": url, "mime": mime}
    context.artifacts.append(artifact)
    if emit:
        await emit(FILE_CARD.build({"title": title, "filename": filename, "id": artifact_id}))
    return {"status": "created", "filename": filename, "download_url": url}


async def create_document_artifact(context, title, sections, emit=None):
    content = render_document(title, sections, context.client["display_name"])
    await reserve_tool(context, "documents")
    return await save_artifact(context, title, ".docx", DOCX_MIME, content, emit)


async def create_image_artifact(context, title, prompt, emit=None):
    if not 1 <= len(title) <= 200 or not 1 <= len(prompt) <= 4000:
        raise image_generation_error("image_invalid_request")
    if not context.config.get("images"):
        raise image_generation_error("image_tool_disabled")
    # Surface the actual failed attempt promptly. ChatKit offers an explicit
    # retry only for transient failures; quota/policy failures are not retried.
    api = openai_hosted_client(context, timeout=75, max_retries=0)
    if api is None:
        raise image_generation_error("image_provider_unconfigured")
    try:
        await reserve_tool(context, "images")
    except ValueError as error:
        code = {
            "At most two files can be generated in one reply": "image_reply_limit_reached",
            "The client's monthly generation allowance has been reached": "image_generation_limit_reached",
        }.get(str(error), "image_generation_failed")
        raise image_generation_error(code) from None
    try:
        # The hosted client may have been reused by a different tool earlier in
        # this reply. Override retry policy on this image call as well.
        image_model = os.getenv("OPENAI_IMAGE_MODEL", "gpt-image-2").strip() or "gpt-image-2"
        if hasattr(context, "cost_events"):
            context.cost_events.append({"provider": "openai", "model": image_model, "kind": "image"})
        result = await api.with_options(timeout=75, max_retries=0).images.generate(
            model=image_model,
            prompt="Approved client instructions:\n" + context.config["instructions"]
                + "\n\nRequested image:\n" + prompt,
            n=1, size="1024x1024", quality="medium", output_format="webp",
        )
        if not result.data or not result.data[0].b64_json:
            raise image_generation_error("image_response_invalid")
        try:
            content = base64.b64decode(result.data[0].b64_json, validate=True)
        except (ValueError, TypeError):
            raise image_generation_error("image_response_invalid") from None
        if not content:
            raise image_generation_error("image_response_invalid")
        return await save_artifact(context, title, ".webp", "image/webp", content, emit)
    except ImageGenerationError:
        raise
    except Exception as error:
        raise image_provider_error(error) from None


def generation_tools(context, emit=None):
    @function_tool(failure_error_function=tool_failure)
    async def create_document(title: str, sections: list[DocumentSection]) -> dict:
        """Create a downloadable Word document from approved content. Use headings, paragraphs and bullets; ask for missing facts first."""
        return await create_document_artifact(context, title, sections, emit)

    # A failed image request must reach ChatKit's server error event, not become
    # model input that can be paraphrased into a fictitious success or download.
    @function_tool(failure_error_function=None)
    async def create_image(title: str, prompt: str) -> dict:
        """Generate one image when the user asks, following approved client instructions and branding. Describe the intended image in prompt."""
        return await create_image_artifact(context, title, prompt, emit)

    image_available = context.config.get("images") and has_openai_hosted_key(context)
    return ([create_document] if context.config.get("documents") else []) + ([create_image] if image_available else [])
